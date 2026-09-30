# -*- coding: utf-8 -*-
"""Create and reopen a small DOCX report without requiring Microsoft Word."""
import argparse
import json
from pathlib import Path
import sys


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
        document = Document()
        style = document.styles["Normal"]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(11)
        properties = style.element.get_or_add_rPr()
        fonts = properties.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            properties.append(fonts)
        fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        document.add_heading("项目报告", level=0)
        document.add_paragraph("这是使用随包 Python 生成的中文示例。")
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "项目"
        table.rows[0].cells[1].text = "结果"
        row = table.add_row().cells
        row[0].text = "结构检查"
        row[1].text = "保存后重新读取"
        output = Path(args.output).expanduser().resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("xb") as handle:
            document.save(handle)
        reopened = Document(str(output))
        if reopened.paragraphs[0].text != "项目报告" or reopened.tables[0].cell(1, 1).text != "保存后重新读取":
            raise ValueError("Saved DOCX content mismatch")
        print(json.dumps({"ok": True, "output": str(output), "paragraphs": len(reopened.paragraphs),
                          "tables": len(reopened.tables), "verified": True,
                          "visualLayoutChecked": False}, ensure_ascii=False))
        return 0
    except Exception as error:
        print(json.dumps({"ok": False, "error": "{}: {}".format(type(error).__name__, error)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    sys.exit(main())
