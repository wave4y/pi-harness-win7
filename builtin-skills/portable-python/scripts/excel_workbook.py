# -*- coding: utf-8 -*-
"""Create an XLSX report with formatting, a formula and its explicit cached result."""
import argparse
import json
from pathlib import Path
import sys


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        import openpyxl
        import xlsxwriter
        output = Path(args.output).expanduser().resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("xb") as handle:
            with xlsxwriter.Workbook(handle) as book:
                sheet = book.add_worksheet("销售")
                header = book.add_format({"bold": True, "bg_color": "#E8EFF8"})
                number = book.add_format({"num_format": "0.00"})
                sheet.write_row("A1", ["项目", "数量", "单价", "金额"], header)
                sheet.write_row("A2", ["中文样例", 3, 12.5])
                sheet.write_formula("D2", "=B2*C2", number, 37.5)
                sheet.set_column("A:A", 22)
                sheet.set_column("B:D", 14)
                sheet.freeze_panes(1, 0)
        formulas = openpyxl.load_workbook(str(output), data_only=False, read_only=True)
        try:
            if formulas["销售"]["D2"].value != "=B2*C2" or formulas["销售"]["A2"].value != "中文样例":
                raise ValueError("Formula or text did not persist")
        finally:
            formulas.close()
        values = openpyxl.load_workbook(str(output), data_only=True, read_only=True)
        try:
            if values["销售"]["D2"].value != 37.5:
                raise ValueError("Cached formula value did not persist")
        finally:
            values.close()
        print(json.dumps({"ok": True, "output": str(output), "verified": True,
                          "formulaRecalculated": False}, ensure_ascii=False))
        return 0
    except Exception as error:
        print(json.dumps({"ok": False, "error": "{}: {}".format(type(error).__name__, error)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    sys.exit(main())
