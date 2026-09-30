# -*- coding: utf-8 -*-
"""Offline portable-Python checks. Only writes within a new output subdirectory."""
import argparse
import importlib
import importlib.metadata
import json
from pathlib import Path
import platform
import sys
import tempfile


def require(condition, message):
    if not condition:
        raise AssertionError(message)


def check_packages(directory):
    packages = [
        ("requests", "requests"), ("numpy", "numpy"), ("pandas", "pandas"),
        ("openpyxl", "openpyxl"), ("XlsxWriter", "xlsxwriter"),
        ("python-docx", "docx"), ("python-pptx", "pptx"),
        ("Pillow", "PIL"), ("lxml", "lxml.etree"),
    ]
    results = []
    for distribution, module in packages:
        try:
            importlib.import_module(module)
            results.append({"package": distribution, "ok": True,
                            "version": importlib.metadata.version(distribution)})
        except Exception as error:
            results.append({"package": distribution, "ok": False,
                            "error": "{}: {}".format(type(error).__name__, error)})
    return {"ok": all(item["ok"] for item in results), "packages": results,
            "networkRequests": 0}


def check_numpy(directory):
    import numpy as np
    matrix = np.array([[3.0, 1.0], [1.0, 2.0]])
    vector = np.array([9.0, 8.0])
    solution = np.linalg.solve(matrix, vector)
    require(np.allclose(solution, [2.0, 3.0]), "NumPy linear solve mismatch")
    require(np.allclose(matrix @ solution, vector), "NumPy matrix product mismatch")
    return {"solution": solution.tolist(), "determinant": float(np.linalg.det(matrix))}


def check_csv(directory):
    import pandas as pd
    path = directory / "中文 表格.csv"
    expected = pd.DataFrame({"项目": ["甲", "乙"], "数量": [2, 3], "金额": [7.5, 12.0]})
    expected.to_csv(str(path), index=False, encoding="utf-8-sig", mode="x")
    actual = pd.read_csv(str(path), encoding="utf-8-sig", dtype={"项目": str})
    pd.testing.assert_frame_equal(actual, expected)
    return {"files": [str(path)], "rows": len(actual), "columns": list(actual.columns)}


def check_excel(directory):
    import openpyxl
    import xlsxwriter
    path = directory / "中文 工作簿.xlsx"
    edited_path = directory / "中文 工作簿 重存.xlsx"
    with path.open("xb") as output:
        with xlsxwriter.Workbook(output) as workbook:
            sheet = workbook.add_worksheet("数据")
            sheet.write_row("A1", ["项目", "数量", "单价", "金额"])
            sheet.write_row("A2", ["中文 样例", 3, 12.5])
            sheet.write_formula("D2", "=B2*C2", None, 37.5)
    formulas = openpyxl.load_workbook(str(path), data_only=False)
    try:
        require(formulas["数据"]["A2"].value == "中文 样例", "XLSX Unicode cell mismatch")
        require(formulas["数据"]["D2"].value == "=B2*C2", "XLSX formula mismatch")
        formulas["数据"]["A3"] = "编辑后重开"
        with edited_path.open("xb") as output:
            formulas.save(output)
    finally:
        formulas.close()
    values = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        require(values["数据"]["D2"].value == 37.5, "XLSX cached formula result mismatch")
    finally:
        values.close()
    edited = openpyxl.load_workbook(str(edited_path), data_only=False, read_only=True)
    try:
        require(edited["数据"]["A3"].value == "编辑后重开", "XLSX edit did not persist")
        require(edited["数据"]["D2"].value == "=B2*C2", "XLSX edit lost formula")
    finally:
        edited.close()
    return {"files": [str(path), str(edited_path)], "worksheet": "数据",
            "cachedFormulaResult": 37.5, "formulaRecalculated": False}


def check_word(directory):
    from docx import Document
    path = directory / "中文 文档.docx"
    document = Document()
    document.add_heading("便携 Python 自检", level=1)
    document.add_paragraph("中文段落与 UTF-8 路径。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "结果"
    table.cell(1, 0).text = "文档重开"
    table.cell(1, 1).text = "通过"
    with path.open("xb") as output:
        document.save(output)
    reopened = Document(str(path))
    require(reopened.paragraphs[1].text == "中文段落与 UTF-8 路径。", "DOCX paragraph mismatch")
    require(reopened.tables[0].cell(1, 1).text == "通过", "DOCX table mismatch")
    return {"files": [str(path)], "paragraphs": len(reopened.paragraphs),
            "tables": len(reopened.tables), "visualLayoutChecked": False}


def check_powerpoint(directory):
    from pptx import Presentation
    from pptx.util import Inches
    path = directory / "中文 幻灯片.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    textbox.text_frame.text = "便携 Python 自检\n中文幻灯片"
    with path.open("xb") as output:
        presentation.save(output)
    reopened = Presentation(str(path))
    require(len(reopened.slides) == 1, "PPTX slide count mismatch")
    texts = [shape.text for shape in reopened.slides[0].shapes if shape.has_text_frame]
    require("便携 Python 自检\n中文幻灯片" in texts, "PPTX text mismatch")
    return {"files": [str(path)], "slides": len(reopened.slides), "visualLayoutChecked": False}


def check_images(directory):
    from PIL import Image, ImageDraw
    png_path = directory / "中文 图片.png"
    jpeg_path = directory / "中文 图片.jpg"
    with Image.new("RGB", (96, 64), (30, 80, 140)) as picture:
        ImageDraw.Draw(picture).rectangle((10, 10, 35, 35), fill=(240, 190, 50))
        with png_path.open("xb") as output:
            picture.save(output, format="PNG")
        with jpeg_path.open("xb") as output:
            picture.save(output, format="JPEG", quality=90)
    results = []
    for path, expected_format in [(png_path, "PNG"), (jpeg_path, "JPEG")]:
        with Image.open(str(path)) as reopened:
            reopened.load()
            require(reopened.size == (96, 64), expected_format + " size mismatch")
            require(reopened.format == expected_format, expected_format + " format mismatch")
            if expected_format == "PNG":
                require(reopened.getpixel((15, 15)) == (240, 190, 50), "PNG pixel mismatch")
            results.append({"format": reopened.format, "size": list(reopened.size)})
    return {"files": [str(png_path), str(jpeg_path)], "images": results}


def check_xml(directory):
    from lxml import etree
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    document = etree.fromstring('<根><结果>通过</结果></根>'.encode("utf-8"), parser)
    require(document.findtext("结果") == "通过", "XML Unicode parse mismatch")
    return {"root": document.tag, "externalEntitiesEnabled": False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", required=True, help="Parent directory for a new unique results folder")
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    report = {"ok": False, "offline": True, "python": {
        "version": platform.python_version(), "executable": sys.executable,
        "platform": platform.platform(), "architecture": platform.machine()}, "checks": []}
    try:
        output_parent = Path(args.output_dir).expanduser().resolve()
        output_parent.mkdir(parents=True, exist_ok=True)
        directory = Path(tempfile.mkdtemp(prefix="Python 自检 ", dir=str(output_parent)))
        report["outputDirectory"] = str(directory)
        checks = [("imports", check_packages), ("numpy", check_numpy), ("csv", check_csv),
                  ("xlsx", check_excel), ("docx", check_word), ("pptx", check_powerpoint),
                  ("images", check_images), ("xml", check_xml)]
        for name, check in checks:
            try:
                result = check(directory)
                result.setdefault("ok", True)
                result["name"] = name
            except Exception as error:
                result = {"name": name, "ok": False,
                          "error": "{}: {}".format(type(error).__name__, error)}
            report["checks"].append(result)
        report["ok"] = all(result["ok"] for result in report["checks"])
    except Exception as error:
        report["error"] = "{}: {}".format(type(error).__name__, error)
    report["limitations"] = ["No network requests were made.",
                             "Office layout and formula recalculation were not tested.",
                             "Success on this OS does not verify Windows 7 compatibility."]
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
