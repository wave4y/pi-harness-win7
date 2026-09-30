# -*- coding: utf-8 -*-
"""Create and reopen a small widescreen PPTX without requiring PowerPoint."""
import argparse
import json
from pathlib import Path
import sys


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt
        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        for title, body in [("项目汇报", "使用便携 Python 生成的中文示例"),
                            ("结果核对", "保存后重新读取页数与文本；视觉排版需另行检查。")]:
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            for text, top, size in [(title, 0.8, 32), (body, 2.0, 22)]:
                box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(2.2))
                paragraph = box.text_frame.paragraphs[0]
                paragraph.text = text
                paragraph.font.name = "Microsoft YaHei"
                paragraph.font.size = Pt(size)
                paragraph.font.color.rgb = RGBColor(35, 51, 72)
                box.text_frame.word_wrap = True
        output = Path(args.output).expanduser().resolve()
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open("xb") as handle:
            presentation.save(handle)
        reopened = Presentation(str(output))
        texts = [shape.text for slide in reopened.slides for shape in slide.shapes if shape.has_text_frame]
        if len(reopened.slides) != 2 or "项目汇报" not in texts or "结果核对" not in texts:
            raise ValueError("Saved PPTX content mismatch")
        print(json.dumps({"ok": True, "output": str(output), "slides": len(reopened.slides),
                          "verified": True, "visualLayoutChecked": False}, ensure_ascii=False))
        return 0
    except Exception as error:
        print(json.dumps({"ok": False, "error": "{}: {}".format(type(error).__name__, error)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    sys.exit(main())
